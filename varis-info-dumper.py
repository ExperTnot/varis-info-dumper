import os
import re
import time
from docx import Document
import openpyxl

CONFIG_FILE_DOCX = "config.txt"
CONFIG_FILE_XLSX = "config_xlsx.txt"

# Function to extract data from a Word document
def extract_word_data(docx_file):
    document = Document(docx_file)
    data = []

    for paragraph in document.paragraphs:
        text = paragraph.text
        occurrences = []
        start = 0

        while True:
            start = text.find("HP:", start)

            if start == -1:
                break

            end = start + 10  # 10 characters (HP:0000000)
            hp_number = text[start:end]
            occurrences.append(hp_number)
            start = end

        if occurrences:
            data.extend(occurrences)

    return data

# Function to open a file with a 1-second delay
def open_file_with_delay(file_path):
    time.sleep(1)
    os.system(f"start {file_path}")

# Function to get the folder path from the user and save it to the configuration file
def get_folder_path():
    folder_path = input("Enter the folder path where your .docx files are located: ")
    with open(CONFIG_FILE_DOCX, "w") as config_file:
        config_file.write(folder_path)
    return folder_path

# Function to read the folder path from the configuration file
def read_folder_path():
    if os.path.exists(CONFIG_FILE_DOCX):
        with open(CONFIG_FILE_DOCX, "r") as config_file:
            return config_file.read()
    else:
        return None

# Function to search for a 4-digit number in an Excel sheet and extract data from a specific column (Column C -> Column I)
def search_excel_and_extract_data(excel_file, search_value):
    wb = openpyxl.load_workbook(excel_file)
    sheet = wb.active

    extracted_data = None

    # Iterate through rows in Column C
    for row in sheet.iter_rows(min_row=2, values_only=True):
        cell_value = str(row[2])  # Column C is the third column (index 2)

        if search_value in cell_value:
            # Extract data from Column I (if available)
            column_i_value = None
            if len(row) > 8:  # Check if there are enough columns in the current row
                column_i_value = row[8]  # Column I (index 8)

            extracted_data = (cell_value, column_i_value)
            break  # Stop iterating after the first match

    wb.close()
    return extracted_data

# Function to get the Excel file path from the user and save it to the configuration file
def get_excel_file_path():
    excel_file_path = input("Enter the path to the Excel sheet (xlsx file): ")
    with open(CONFIG_FILE_XLSX, "w") as config_file:
        config_file.write(excel_file_path)
    return excel_file_path

# Function to read the Excel file path from the configuration file
def read_excel_file_path():
    if os.path.exists(CONFIG_FILE_XLSX):
        with open(CONFIG_FILE_XLSX, "r") as config_file:
            return config_file.read()
    else:
        return None

# Function to add data to a text file
def add_data_to_text_file(file_path, data):
    with open(file_path, "a") as file:
        if data is not None:
            file.write(data + "\n")  # Append data with a newline character

def main():
    # Check if the configuration file for Word documents exists
    folder_path = read_folder_path()

    if not folder_path:
        print("Configuration file for Word documents not found or folder path not configured.")
        folder_path = get_folder_path()

    # Check if the configuration file for Excel files exists
    excel_file_path = read_excel_file_path()

    if not excel_file_path:
        print("Configuration file for Excel files not found or Excel file path not configured.")
        excel_file_path = get_excel_file_path()

    if not os.path.exists(excel_file_path):
        print(f"Excel file '{excel_file_path}' not found.")
        return

    # Input: Provide a partial number to search for
    partial_number = input("Enter a partial number to search for: ")

    folder_dir = folder_path

    matching_folders = []
    for folder_name in os.listdir(folder_dir):
        if os.path.isdir(os.path.join(folder_dir, folder_name)) and folder_name.startswith(partial_number):
            matching_folders.append(folder_name)

    if not matching_folders:
        print(f"No matching folders found for '{partial_number}' in Word documents.")
        return

    script_dir = os.path.dirname(__file__)

    for folder_name in matching_folders:
        folder_path = os.path.join(folder_dir, folder_name)
        docx_files = [file for file in os.listdir(folder_path) if file.endswith(".docx") and file.startswith(folder_name)]

        if not docx_files:
            print(f"No .docx files with the same leading number found in folder '{folder_name}'.")
            continue

        docx_file = os.path.join(folder_path, docx_files[0])  # Use the first .docx file found

        word_data = extract_word_data(docx_file)

        # Path for the output text file in the same directory as the script
        output_file_path = os.path.join(script_dir, f"{folder_name}.txt")

        # Write the collected HP numbers to the output text file
        with open(output_file_path, "w") as output_file:
            for hp_number in word_data:
                output_file.write(hp_number + "\n")

        print(f"Data from {docx_files[0]} has been saved to {folder_name}.txt")

        # Open the generated .txt file with a 1-second delay
        open_file_with_delay(output_file_path)

    # Input: Provide a 4-digit number to search for in the Excel sheet (xlsx file)
    search_value = partial_number #input("Enter a 4-digit number to search for in the Excel sheet: ")

    # Search the Excel sheet and extract data
    extracted_data = search_excel_and_extract_data(excel_file_path, search_value)

    if extracted_data is None:
        print(f"No data found for '{search_value}' in the Excel sheet.")
        return

    # Display the extracted data
    cell_value, data = extracted_data
    print(f"Cell Value: {cell_value}")
    if data is not None:
        print(f"Data: {data}\n")
    else:
        print("No data found in Column I\n")

    # Create a text file with the cell value as the filename and add the data to the file
    text_file_path = os.path.join(script_dir, f"{cell_value}.txt")
    add_data_to_text_file(text_file_path, data)

if __name__ == "__main__":
    main()
